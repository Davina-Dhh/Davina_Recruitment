# -*- coding: utf-8 -*-
"""简历文件解析：PDF / Word（.docx）→ 纯文本"""
from __future__ import annotations

import io
import re
from typing import Tuple


def extract_resume_text(filename: str, raw: bytes) -> Tuple[str, str]:
    """
    返回 (text, format_label)。
    支持 .pdf / .docx；.doc 旧格式不支持。
    """
    name = (filename or "resume").strip()
    lower = name.lower()
    if not raw:
        raise ValueError("文件为空")

    if lower.endswith(".pdf"):
        text = _from_pdf(raw)
        return _clean(text), "PDF"
    if lower.endswith(".docx"):
        text = _from_docx(raw)
        return _clean(text), "Word"
    if lower.endswith(".doc"):
        raise ValueError("暂不支持旧版 .doc，请另存为 .docx 或导出 PDF")
    raise ValueError("仅支持 PDF 或 Word（.docx）")


def _from_pdf(raw: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise RuntimeError("缺少 pypdf，请执行: pip install pypdf") from e
    reader = PdfReader(io.BytesIO(raw))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    text = "\n".join(parts).strip()
    if len(text) < 20:
        raise ValueError("未能从 PDF 提取到有效文字（可能是扫描件图片，请用可选中文字的 PDF）")
    return text


def _from_docx(raw: bytes) -> str:
    try:
        from docx import Document
    except ImportError as e:
        raise RuntimeError("缺少 python-docx，请执行: pip install python-docx") from e
    doc = Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs if (p.text or "").strip()]
    # 表格文字
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if len(text) < 20:
        raise ValueError("未能从 Word 提取到有效文字")
    return text


def _clean(text: str) -> str:
    t = text.replace("\x00", " ")
    t = re.sub(r"[ \t]+\n", "\n", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()
